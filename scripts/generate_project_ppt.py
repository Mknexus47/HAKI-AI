from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT_FILE = r"c:\Users\user\Desktop\HAKI AI\HAKI-AI-Project-Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(12, 16, 24)
DARK = RGBColor(18, 26, 36)
GOLD = RGBColor(201, 162, 39)
GOLD_SOFT = RGBColor(246, 214, 130)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(196, 204, 212)
GREEN = RGBColor(42, 175, 110)
RED = RGBColor(210, 66, 66)



def add_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(9.5), Inches(0.7))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.name = 'Aptos'
    run.font.color.rgb = WHITE

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.72), Inches(1.0), Inches(9.5), Inches(0.35))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.runs[0]
        run2.font.size = Pt(11)
        run2.font.name = 'Aptos'
        run2.font.color.rgb = MUTED



def add_header_line(slide):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.45), Inches(11.8), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()


def add_footer(slide, label):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(11.8), Inches(0.25))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.RIGHT
    run = p.runs[0]
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    run.font.name = 'Aptos'


# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)

# Brand accent block
brand = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.5), Inches(0.3), Inches(0.3))
brand.fill.solid(); brand.fill.fore_color.rgb = GOLD
brand.line.fill.background()
brand.shadow.inherit = False

# Title text
box = slide.shapes.add_textbox(Inches(1.1), Inches(1.0), Inches(7.2), Inches(1.5))
tf = box.text_frame
p = tf.paragraphs[0]
p.text = 'HAKI AI'
p.alignment = PP_ALIGN.LEFT
run = p.runs[0]
run.font.size = Pt(30)
run.font.bold = True
run.font.name = 'Aptos'
run.font.color.rgb = WHITE

sub = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(8.5), Inches(0.7))
tf2 = sub.text_frame
p2 = tf2.paragraphs[0]
p2.text = 'AI-powered legal guidance, document generation and dispute support for everyday Kenyan users.'
p2.alignment = PP_ALIGN.LEFT
run2 = p2.runs[0]
run2.font.size = Pt(18)
run2.font.name = 'Aptos'
run2.font.color.rgb = MUTED

# Stats cards
cards = [
    ('Legal topics', '15+ practical legal areas'),
    ('AI assistant', 'Plain-language answers'),
    ('Document tools', 'Letters, affidavits, agreements'),
    ('User trust', 'General-info + escalation paths'),
]
for i, (title, body) in enumerate(cards):
    x = Inches(0.9 + i * 2.8)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.1), Inches(2.5), Inches(1.4))
    card.fill.solid(); card.fill.fore_color.rgb = NAVY
    card.line.color.rgb = GOLD
    card.line.width = Pt(1.2)
    tfc = card.text_frame
    tfc.word_wrap = True
    p = tfc.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.name = 'Aptos'
    run.font.color.rgb = WHITE
    p2 = tfc.add_paragraph()
    p2.text = body
    p2.alignment = PP_ALIGN.LEFT
    run2 = p2.runs[0]
    run2.font.size = Pt(10)
    run2.font.name = 'Aptos'
    run2.font.color.rgb = MUTED

add_footer(slide, 'Project overview | HAKI AI')

# Slide 2: Problem
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, 'The problem we are solving', 'Access to legal help is often slow, expensive, and confusing for everyday people.')
add_header_line(slide)

left = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.2), Inches(4.1))
tf = left.text_frame
for idx, text in enumerate([
    '• Most Kenyans cannot afford a lawyer for routine legal questions.',
    '• Legal information is often buried in formal language and scattered sources.',
    '• People need guidance before disputes escalate or documents are drafted incorrectly.',
    '• A simple first step is missing for many everyday legal situations.',
]):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.text = text
    p.level = 0
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(22 if idx < 2 else 18)
    run.font.name = 'Aptos'
    run.font.color.rgb = WHITE

right = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.2), Inches(2.0), Inches(4.8), Inches(3.6))
right.fill.solid(); right.fill.fore_color.rgb = DARK
right.line.color.rgb = GOLD
right.line.width = Pt(1.5)
rt = right.text_frame
rt.word_wrap = True
p = rt.paragraphs[0]
p.text = 'Why this matters'
p.alignment = PP_ALIGN.CENTER
run = p.runs[0]
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = GOLD_SOFT
for t in ['Delay in action','Wrong document drafting','Stress during disputes','Loss of rights or deposits']:
    p2 = rt.add_paragraph(); p2.text = t; p2.alignment = PP_ALIGN.LEFT; run2 = p2.runs[0]; run2.font.size = Pt(15); run2.font.color.rgb = WHITE
add_footer(slide, 'Problem statement')

# Slide 3: Product overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_title(slide, 'What HAKI AI does', 'A conversational legal assistant with practical document generation and trust-first guidance.')
add_header_line(slide)

features = [
    ('Ask a legal question', 'Get plain-language explanations on rights, obligations, contracts, disputes and procedures.'),
    ('Generate legal documents', 'Create letters, acknowledgements, affidavits and business-ready agreement drafts.'),
    ('Explore legal topics', 'Browse common topic clusters such as tenancy, employment, consumer issues and debt matters.'),
    ('Get structured next steps', 'Receive step-by-step guidance, escalation paths and risk-aware recommendations.'),
]
for i, (title, body) in enumerate(features):
    x = Inches(0.7 + (i % 2) * 6.0)
    y = Inches(1.8 + (i // 2) * 2.1)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.3), Inches(1.6))
    card.fill.solid(); card.fill.fore_color.rgb = NAVY
    card.line.color.rgb = GOLD
    card.line.width = Pt(1.2)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.text = body
    p2.alignment = PP_ALIGN.LEFT
    run2 = p2.runs[0]
    run2.font.size = Pt(11)
    run2.font.color.rgb = MUTED
add_footer(slide, 'Product overview')

# Slide 4: Key features
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, 'Platform features', 'Purpose-built to turn legal uncertainty into quick, usable next steps.')
add_header_line(slide)

feature_rows = [
    ('AI legal assistant', 'Answers user questions in clear, everyday language and explains rights and next actions.'),
    ('Document generator', 'Drafts templates for common legal letters and acknowledgements in a few steps.'),
    ('RAG-based knowledge retrieval', 'Pulls from curated legal sources and topic data to improve grounded responses.'),
    ('Trust and safety', 'Adds risk disclaimers, general information boundaries, and escalation guidance.'),
]
for i, (title, desc) in enumerate(feature_rows):
    y = Inches(1.8 + i * 1.35)
    tag = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y, Inches(0.18), Inches(0.18))
    tag.fill.solid(); tag.fill.fore_color.rgb = GOLD
    tag.line.fill.background()
    txt = slide.shapes.add_textbox(Inches(1.2), y - 0.06, Inches(10.8), Inches(0.55))
    tf = txt.text_frame
    p = tf.paragraphs[0]
    p.text = f'{title} — {desc}'
    run = p.runs[0]
    run.font.size = Pt(20)
    run.font.bold = True if i in [0,2] else False
    run.font.color.rgb = WHITE
    run.font.name = 'Aptos'
add_footer(slide, 'Core features')

# Slide 5: Target users and use cases
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_title(slide, 'Who it is for', 'Built for people who need reliable first-step legal help without friction.')
add_header_line(slide)

users = [
    'Tenants and landlords dealing with rent, notices and deposits',
    'Employees seeking clarity on contracts, wages or termination issues',
    'Consumers facing debt, refund, service or complaint disputes',
    'Small businesses drafting simple legal documents and agreement basics',
]
for i, item in enumerate(users):
    x = Inches(0.9)
    y = Inches(1.9 + i * 1.35)
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.3), Inches(0.8))
    c.fill.solid(); c.fill.fore_color.rgb = NAVY
    c.line.color.rgb = GOLD
    tf = c.text_frame
    p = tf.paragraphs[0]
    p.text = item
    run = p.runs[0]
    run.font.size = Pt(18)
    run.font.color.rgb = WHITE

journey = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.1), Inches(2.1), Inches(4.9), Inches(3.6))
journey.fill.solid(); journey.fill.fore_color.rgb = NAVY
journey.line.color.rgb = GOLD
jt = journey.text_frame
jt.word_wrap = True
steps = ['1. Ask a question', '2. Get plain-language guidance', '3. Review the suggested document', '4. Take action or escalate to a professional']
for idx, step in enumerate(steps):
    p = jt.paragraphs[0] if idx == 0 else jt.add_paragraph()
    p.text = step
    run = p.runs[0]
    run.font.size = Pt(17 if idx < 2 else 16)
    run.font.color.rgb = WHITE
    if idx == 0:
        p.level = 0
add_footer(slide, 'Use cases')

# Slide 6: Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, 'Technical architecture', 'A practical stack combining AI, retrieval, and legal-document workflows.')
add_header_line(slide)

layers = [
    ('Frontend', 'Next.js app with legal topic pages, chat flows, and document generation UI'),
    ('Application layer', 'TypeScript services for chat, moderation, RAG orchestration, and legal retrieval'),
    ('Data layer', 'Supabase + SQL data, topic libraries, and document-related metadata'),
    ('AI layer', 'Model providers, grounding, retrieval, and legal response generation'),
]
for i, (name, desc) in enumerate(layers):
    x = Inches(0.7 + (i % 2) * 5.9)
    y = Inches(1.9 + (i // 2) * 2.0)
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.2), Inches(1.4))
    c.fill.solid(); c.fill.fore_color.rgb = DARK
    c.line.color.rgb = GOLD
    tf = c.text_frame
    p = tf.paragraphs[0]
    p.text = name
    run = p.runs[0]
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = GOLD_SOFT
    p2 = tf.add_paragraph(); p2.text = desc; run2 = p2.runs[0]; run2.font.size = Pt(11); run2.font.color.rgb = WHITE

add_footer(slide, 'Architecture overview')

# Slide 7: RAG and legal knowledge
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_title(slide, 'Legal knowledge and grounding', 'The product is designed to improve clarity, relevance and safer legal assistance.')
add_header_line(slide)

left = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.4))
tf = left.text_frame
for idx, text in enumerate([
    '• Curated legal topics and public-source material provide grounding for common scenarios.',
    '• Retrieval complements generation with more relevant snippets and topic context.',
    '• The system is designed to explain legal concepts in accessible language, not just produce raw text.',
    '• Responses can direct users to next steps, documentation, or professional support when needed.',
]):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.text = text
    run = p.runs[0]
    run.font.size = Pt(20)
    run.font.color.rgb = WHITE
    run.font.name = 'Aptos'

badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(2.0), Inches(3.8), Inches(2.2))
badge.fill.solid(); badge.fill.fore_color.rgb = NAVY
badge.line.color.rgb = GOLD
b = badge.text_frame
p = b.paragraphs[0]
p.text = 'Grounding principles'
p.alignment = PP_ALIGN.CENTER
run = p.runs[0]
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = GOLD_SOFT
for line in ['Accuracy', 'Clarity', 'Context', 'Caution']:
    p2 = b.add_paragraph(); p2.text = line; p2.alignment = PP_ALIGN.CENTER; run2 = p2.runs[0]; run2.font.size = Pt(18); run2.font.color.rgb = WHITE
add_footer(slide, 'Knowledge layer')

# Slide 8: Business model
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, 'Business model and value', 'The product creates clear value for both everyday users and the legal ecosystem.')
add_header_line(slide)

cards = [
    ('Users', 'Affordable first-step legal support for a wide audience'),
    ('Businesses', 'Document drafting and guidance for SMEs and service providers'),
    ('Partners', 'Potential legal education, referral and distribution opportunities'),
    ('Revenue', 'Freemium access, subscriptions, premium templates, and professional referral channels'),
]
for i, (title, body) in enumerate(cards):
    x = Inches(0.7 + (i % 2) * 5.9)
    y = Inches(1.9 + (i // 2) * 2.0)
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.2), Inches(1.4))
    c.fill.solid(); c.fill.fore_color.rgb = DARK
    c.line.color.rgb = GOLD
    tf = c.text_frame
    p = tf.paragraphs[0]; p.text = title; run = p.runs[0]; run.font.size = Pt(17); run.font.bold = True; run.font.color.rgb = GOLD_SOFT
    p2 = tf.add_paragraph(); p2.text = body; run2 = p2.runs[0]; run2.font.size = Pt(11); run2.font.color.rgb = WHITE
add_footer(slide, 'Value proposition')

# Slide 9: Roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_title(slide, 'Roadmap', 'A focused product path from MVP to trusted legal assistant.')
add_header_line(slide)

roadmap = [
    ('Phase 1', 'Launch legal assistant and core topic coverage for tenancy, employment and consumer issues'),
    ('Phase 2', 'Expand document generation and workflow templates for common legal letters and acknowledgements'),
    ('Phase 3', 'Add stronger data quality controls, legal-source curation and safer escalation guidance'),
    ('Phase 4', 'Scale to partnerships, premium workflows, and wider legal education experiences'),
]
for i, (phase, desc) in enumerate(roadmap):
    y = Inches(1.8 + i * 1.4)
    c = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y, Inches(1.2), Inches(0.7))
    c.fill.solid(); c.fill.fore_color.rgb = GOLD
    c.line.fill.background()
    tf = c.text_frame; p = tf.paragraphs[0]; p.text = phase; run = p.runs[0]; run.font.size = Pt(14); run.font.bold = True; run.font.color.rgb = NAVY
    txt = slide.shapes.add_textbox(Inches(2.3), y, Inches(9.3), Inches(0.8))
    tf2 = txt.text_frame; p2 = tf2.paragraphs[0]; p2.text = desc; run2 = p2.runs[0]; run2.font.size = Pt(18); run2.font.color.rgb = WHITE
add_footer(slide, 'Go-to-market roadmap')

# Slide 10: Closing
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, 'Conclusion', 'HAKI AI is positioned to make legal help more approachable, faster and more practical for everyday users.')
add_header_line(slide)

callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(11.2), Inches(2.4))
callout.fill.solid(); callout.fill.fore_color.rgb = DARK
callout.line.color.rgb = GOLD
cf = callout.text_frame
for idx, text in enumerate([
    'Make legal guidance more accessible',
    'Turn common legal questions into actionable next steps',
    'Support users with trusted, human-friendly document workflows',
]):
    p = cf.paragraphs[0] if idx == 0 else cf.add_paragraph()
    p.text = text
    run = p.runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = WHITE

box = slide.shapes.add_textbox(Inches(1.1), Inches(4.8), Inches(10.8), Inches(1.0))
tf = box.text_frame
p = tf.paragraphs[0]
p.text = 'The opportunity is to create a practical legal AI experience that is useful, trustworthy and locally relevant.'
run = p.runs[0]
run.font.size = Pt(22)
run.font.color.rgb = MUTED
add_footer(slide, 'Thank you')

prs.save(OUT_FILE)
print(f'Created presentation: {OUT_FILE}')
