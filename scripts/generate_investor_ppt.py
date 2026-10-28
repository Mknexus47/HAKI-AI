from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT_FILE = r"c:\Users\user\Desktop\HAKI AI\HAKI-AI-Investor-Deck.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(9, 18, 33)
NAVY_2 = RGBColor(16, 30, 46)
GOLD = RGBColor(197, 160, 63)
GOLD_LIGHT = RGBColor(239, 214, 146)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(203, 211, 220)
GRAY = RGBColor(65, 75, 89)
GREEN = RGBColor(62, 164, 88)


def bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def title(slide, t, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(9.5), Inches(0.8))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = t
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.name = 'Aptos'
    run.font.color.rgb = WHITE

    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.72), Inches(1.1), Inches(9.5), Inches(0.4))
        tf2 = box2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.runs[0]
        run2.font.size = Pt(11)
        run2.font.name = 'Aptos'
        run2.font.color.rgb = MUTED

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.58), Inches(11.7), Inches(0.04))
    line.fill.solid(); line.fill.fore_color.rgb = GOLD
    line.line.fill.background()


def footer(slide, text):
    box = slide.shapes.add_textbox(Inches(9.7), Inches(7.0), Inches(2.7), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    run = p.runs[0]
    run.font.size = Pt(9)
    run.font.name = 'Aptos'
    run.font.color.rgb = MUTED


def add_card(slide, x, y, w, h, title, body, title_color=WHITE, body_color=MUTED):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid(); card.fill.fore_color.rgb = NAVY_2
    card.line.color.rgb = GOLD
    card.line.width = Pt(1.2)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.name = 'Aptos'
    run.font.color.rgb = title_color
    p2 = tf.add_paragraph()
    p2.text = body
    p2.alignment = PP_ALIGN.LEFT
    run2 = p2.runs[0]
    run2.font.size = Pt(10.5)
    run2.font.name = 'Aptos'
    run2.font.color.rgb = body_color


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)

badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.55), Inches(0.34), Inches(0.34))
badge.fill.solid(); badge.fill.fore_color.rgb = GOLD
badge.line.fill.background()

brand = slide.shapes.add_textbox(Inches(1.2), Inches(0.52), Inches(4.0), Inches(0.5))
tf = brand.text_frame
p = tf.paragraphs[0]
p.text = 'HAKI AI'
run = p.runs[0]
run.font.size = Pt(22)
run.font.bold = True
run.font.name = 'Aptos'
run.font.color.rgb = WHITE

main = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.1), Inches(2.0))
tf2 = main.text_frame
p2 = tf2.paragraphs[0]
p2.text = 'AI-powered legal help for everyday Kenyans.'
run2 = p2.runs[0]
run2.font.size = Pt(30)
run2.font.bold = True
run2.font.name = 'Aptos'
run2.font.color.rgb = WHITE

sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(8.0), Inches(1.0))
tf3 = sub.text_frame
p3 = tf3.paragraphs[0]
p3.text = 'A faster path to legal guidance, document drafting, and next-step support for tenants, workers, consumers, and small businesses.'
run3 = p3.runs[0]
run3.font.size = Pt(17)
run3.font.name = 'Aptos'
run3.font.color.rgb = MUTED

metrics = [
    ('15+', 'legal topics'),
    ('AI', 'assistant'),
    ('Docs', 'generated'),
    ('Kenya', 'first focus'),
]
for i, (v, label) in enumerate(metrics):
    x = Inches(0.8 + i * 2.7)
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(4.8), Inches(2.1), Inches(1.2))
    card.fill.solid(); card.fill.fore_color.rgb = NAVY_2
    card.line.color.rgb = GOLD
    tfc = card.text_frame
    tfc.word_wrap = True
    p4 = tfc.paragraphs[0]
    p4.text = v
    p4.alignment = PP_ALIGN.CENTER
    run4 = p4.runs[0]
    run4.font.size = Pt(22)
    run4.font.bold = True
    run4.font.color.rgb = GOLD_LIGHT
    p5 = tfc.add_paragraph(); p5.text = label; p5.alignment = PP_ALIGN.CENTER; run5 = p5.runs[0]; run5.font.size = Pt(10); run5.font.color.rgb = WHITE

tri = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.0))
tri.fill.solid(); tri.fill.fore_color.rgb = NAVY_2
tri.line.color.rgb = GOLD
tri.line.width = Pt(1.5)
tri_tf = tri.text_frame
tri_tf.word_wrap = True
for idx, line in enumerate(['Problem', 'Costly legal uncertainty', 'Lack of simple guidance', 'No reliable first step']):
    p6 = tri_tf.paragraphs[0] if idx == 0 else tri_tf.add_paragraph()
    p6.text = line
    run6 = p6.runs[0]
    run6.font.size = Pt(18 if idx == 0 else 15)
    run6.font.bold = idx == 0
    run6.font.color.rgb = GOLD_LIGHT if idx == 0 else WHITE

footer(slide, 'HAKI AI | Project overview')

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'The problem', 'Legal support is out of reach for most people')

add_card(slide, Inches(0.8), Inches(2.0), Inches(5.2), Inches(1.5), 'Access gap', 'Most people cannot afford a lawyer for routine legal questions and document drafting.')
add_card(slide, Inches(6.3), Inches(2.0), Inches(5.8), Inches(1.5), 'Information gap', 'Legal language is complex, scattered, and difficult for non-specialists to act on.')
add_card(slide, Inches(0.8), Inches(4.1), Inches(5.2), Inches(1.7), 'Risk gap', 'Without a clear first step, small disputes escalate into bigger losses, stress, and missed rights.')
add_card(slide, Inches(6.3), Inches(4.1), Inches(5.8), Inches(1.7), 'Opportunity', 'A trustworthy, accessible legal AI tool can shorten the path from confusion to action.')
footer(slide, 'Problem statement')

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'Our solution', 'A legal AI assistant with practical outputs for real-world situations')

left = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.8), Inches(2.8))
ltf = left.text_frame
for idx, line in enumerate([
    '• Answers common legal questions in simple language.',
    '• Generates standard legal documents and draft letters.',
    '• Suggests next steps, risks, and escalation paths.',
    '• Helps users act before disputes worsen.',
]):
    p = ltf.paragraphs[0] if idx == 0 else ltf.add_paragraph()
    p.text = line
    run = p.runs[0]
    run.font.size = Pt(20)
    run.font.color.rgb = WHITE

right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(2.0), Inches(5.2), Inches(3.4))
right.fill.solid(); right.fill.fore_color.rgb = NAVY_2
right.line.color.rgb = GOLD
rtf = right.text_frame
rtf.word_wrap = True
for idx, line in enumerate(['User asks a legal question', 'AI retrieves relevant legal context', 'System generates a clear answer + document', 'User gets practical next steps']):
    p = rtf.paragraphs[0] if idx == 0 else rtf.add_paragraph()
    p.text = line
    run = p.runs[0]
    run.font.size = Pt(17)
    run.font.color.rgb = WHITE if idx != 0 else GOLD_LIGHT
    run.font.bold = idx == 0

footer(slide, 'Solution overview')

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'Product features', 'Built to be useful, accessible, and trust-aware')

features = [
    ('Legal assistant', 'Plain-language responses across tenancy, employment, consumer, debt, and complaint scenarios.'),
    ('Document generator', 'Drafts common legal letters, acknowledgements, and business agreement templates.'),
    ('Topic explorer', 'Lets users browse core legal topics and understand their rights quickly.'),
    ('Guided actions', 'Shows step-by-step actions, warnings, and escalation options for better decision-making.'),
]
for i, (title_text, body_text) in enumerate(features):
    x = Inches(0.8 + (i % 2) * 5.9)
    y = Inches(2.0 + (i // 2) * 2.0)
    add_card(slide, x, y, Inches(5.2), Inches(1.5), title_text, body_text)
footer(slide, 'Core platform capabilities')

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'How it works', 'A simple flow from user need to useful legal action')

step_specs = [
    ('1. Ask', 'The user enters a legal question or request.'),
    ('2. Context', 'The system retrieves relevant legal topic and source material.'),
    ('3. Response', 'The AI explains the issue in accessible language.'),
    ('4. Action', 'The user can generate a document or follow recommended next steps.'),
]
for i, (step, desc) in enumerate(step_specs):
    x = Inches(0.9 + i * 2.9)
    c = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, Inches(2.2), Inches(2.2), Inches(1.8))
    c.fill.solid(); c.fill.fore_color.rgb = NAVY_2
    c.line.color.rgb = GOLD
    tf = c.text_frame
    p = tf.paragraphs[0]
    p.text = step
    run = p.runs[0]
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = GOLD_LIGHT
    p2 = tf.add_paragraph(); p2.text = desc; run2 = p2.runs[0]; run2.font.size = Pt(9.5); run2.font.color.rgb = WHITE

footer(slide, 'Operational flow')

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'Technical architecture', 'Built on a practical AI + retrieval + app stack')

layers = [
    ('Frontend', 'Next.js web app with pages for chat, topics, and document generation.'),
    ('Application layer', 'TypeScript services for RAG, chat, moderation and legal-source orchestration.'),
    ('Data platform', 'Supabase, SQL datasets, and structured legal topics for local context.'),
    ('AI layer', 'Model providers, retrieval, and grounded answer generation for legal tasks.'),
]
for i, (name, desc) in enumerate(layers):
    x = Inches(0.9 + (i % 2) * 5.8)
    y = Inches(2.0 + (i // 2) * 2.0)
    add_card(slide, x, y, Inches(5.1), Inches(1.5), name, desc, title_color=GOLD_LIGHT)
footer(slide, 'Architecture overview')

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'Why this matters in Kenya', 'A locally relevant product for practical legal access')

add_card(slide, Inches(0.8), Inches(2.0), Inches(3.7), Inches(2.3), 'Tenants', 'Rent disputes, notices, deposits, and habitability issues are common pain points.')
add_card(slide, Inches(4.9), Inches(2.0), Inches(3.7), Inches(2.3), 'Workers', 'Employment issues often require quick, understandable guidance before escalation.')
add_card(slide, Inches(9.0), Inches(2.0), Inches(3.5), Inches(2.3), 'Consumers', 'Refunds, contracts, and complaint processes need clear first steps.')

box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.9), Inches(4.8), Inches(7.5), Inches(1.1))
box.fill.solid(); box.fill.fore_color.rgb = NAVY_2
box.line.color.rgb = GOLD
b_tf = box.text_frame
p = b_tf.paragraphs[0]
p.text = 'The opportunity is to make legal guidance more affordable, faster, and easier to understand.'
run = p.runs[0]; run.font.size = Pt(20); run.font.color.rgb = WHITE
footer(slide, 'Market relevance')

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'Business model', 'A scalable product with multiple value streams')

add_card(slide, Inches(0.8), Inches(2.0), Inches(5.2), Inches(1.7), 'Freemium access', 'Basic legal guidance and lightweight document support for a broad user base.')
add_card(slide, Inches(6.5), Inches(2.0), Inches(5.9), Inches(1.7), 'Premium workflows', 'Advanced document generation, saved drafts, and richer legal templates.')
add_card(slide, Inches(0.8), Inches(4.1), Inches(5.2), Inches(1.8), 'Partnerships', 'Possible integrations with legal education, community organizations, and referral channels.')
add_card(slide, Inches(6.5), Inches(4.1), Inches(5.9), Inches(1.8), 'Trust & safety', 'Clear disclaimers and escalation paths help build credibility and reduce misuse risk.')
footer(slide, 'Value proposition')

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'Roadmap', 'Focused launch path to growth')

roadmap = [
    ('Phase 1', 'Launch core legal assistant and topic library'),
    ('Phase 2', 'Expand document generation workflows and templates'),
    ('Phase 3', 'Increase source quality, trust controls, and fine-tuning'),
    ('Phase 4', 'Scale into partnerships, premium features, and broader legal education'),
]
for i, (phase, desc) in enumerate(roadmap):
    y = Inches(2.0 + i * 1.4)
    pill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y, Inches(1.4), Inches(0.65))
    pill.fill.solid(); pill.fill.fore_color.rgb = GOLD
    pill.line.fill.background()
    pbox = pill.text_frame; p = pbox.paragraphs[0]; p.text = phase; run = p.runs[0]; run.font.size = Pt(14); run.font.bold = True; run.font.color.rgb = NAVY
    txt = slide.shapes.add_textbox(Inches(2.5), y, Inches(9.5), Inches(0.7))
    t_tf = txt.text_frame; p2 = t_tf.paragraphs[0]; p2.text = desc; run2 = p2.runs[0]; run2.font.size = Pt(19); run2.font.color.rgb = WHITE
footer(slide, 'Growth roadmap')

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
title(slide, 'Closing', 'HAKI AI brings legal help closer to the people who need it most')

box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.25), Inches(2.2), Inches(10.8), Inches(2.3))
box.fill.solid(); box.fill.fore_color.rgb = NAVY_2
box.line.color.rgb = GOLD
b_tf = box.text_frame
for idx, line in enumerate([
    'Accessible legal guidance',
    'AI-powered document support',
    'Actionable next steps for everyday users',
]):
    p = b_tf.paragraphs[0] if idx == 0 else b_tf.add_paragraph()
    p.text = line
    run = p.runs[0]
    run.font.size = Pt(24 if idx == 1 else 22)
    run.font.bold = idx == 1
    run.font.color.rgb = WHITE

final_text = slide.shapes.add_textbox(Inches(1.3), Inches(5.2), Inches(10.5), Inches(0.8))
ft = final_text.text_frame
p = ft.paragraphs[0]
p.text = 'We are building a practical, trustworthy legal AI product with clear value for Kenya and beyond.'
run = p.runs[0]; run.font.size = Pt(20); run.font.color.rgb = MUTED
footer(slide, 'Thank you')

prs.save(OUT_FILE)
print(f'Created investor deck: {OUT_FILE}')
